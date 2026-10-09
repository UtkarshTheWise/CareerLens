"""Stage 1 (docs/PIPELINE.md): uploaded file -> text, and PII stripping before any LLM call.

No LLM and no network here. Resume text and file bytes are never logged.
"""

import io
import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field, replace
from typing import Any

from pydantic import BaseModel

from app.errors import ApiError

logger = logging.getLogger("careerlens.ingest")

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 25
MAX_DOCX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 60_000  # a resume typed in the app; far above a few pages


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    page_count: int | None  # None for DOCX and plain text: neither stores a page count


def _plain_text(filename: str | None, data: bytes) -> str | None:
    """A resume built in the app, sent as resume.txt. Only a *.txt name, strict UTF-8, no control characters."""
    if not (filename or "").lower().endswith(".txt"):
        return None
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return None
    if any(unicodedata.category(c) == "Cc" and c not in "\t\n\r" for c in text):
        return None
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _detect_type(filename: str | None, data: bytes) -> str | None:
    """Magic bytes first; plain text only when the file is named *.txt."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
    if _plain_text(filename, data) is not None:
        return "text"
    return None


def _pdf_text(data: bytes) -> ExtractedDocument:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        total = len(pdf.pages)
        if total > MAX_PDF_PAGES:  # a resume is a few pages; thousands would tie up a request thread
            raise ApiError(422, "unsupported_file", f"This PDF has {total} pages; a resume should be a few.")
        pages = [page.extract_text() or "" for page in pdf.pages]
    return ExtractedDocument("\n\n".join(p.strip() for p in pages if p.strip()), len(pages))


def _docx_text(data: bytes) -> ExtractedDocument:
    import docx

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:  # refuse a zip bomb before parsing it
            if sum(i.file_size for i in archive.infolist()) > MAX_DOCX_UNPACKED_BYTES:
                raise ApiError(422, "unsupported_file", "This document is too large once unpacked.")
    except zipfile.BadZipFile:
        pass  # python-docx reports the damaged file below
    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ExtractedDocument("\n".join(line for line in lines if line.strip()), None)


def extract_text(filename: str | None, data: bytes) -> ExtractedDocument:
    """PDF, DOCX or plain-text bytes -> text. Raises ApiError (contract Error shape) on anything unusable."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ApiError(413, "payload_too_large", "File is larger than 5 MB")
    kind = _detect_type(filename, data)
    if kind is None:
        raise ApiError(422, "unsupported_file", "Upload a PDF, DOCX or plain-text (.txt) file")
    try:
        if kind == "text":
            text = _plain_text(filename, data) or ""
            if len(text) > MAX_TEXT_CHARS:
                raise ApiError(422, "unsupported_file", "This text resume is too long; a resume is a few pages.")
            doc = ExtractedDocument(text, None)
        else:
            doc = _pdf_text(data) if kind == "pdf" else _docx_text(data)
    except ApiError:
        raise
    except Exception as exc:  # parser libraries raise many unrelated types on damaged files
        logger.warning("Could not parse %s upload (%d bytes): %s", kind, len(data), type(exc).__name__)
        raise ApiError(
            422, "no_text_extracted", "This file could not be read. Try exporting it again."
        ) from exc
    if len(doc.text.strip()) < 30:
        raise ApiError(
            422,
            "no_text_extracted",
            "No text found in this file. If it is a scanned image, upload a text-based PDF or DOCX instead.",
        )
    doc = replace(doc, text=doc.text.replace("\x00", ""))  # Postgres text columns reject NUL characters
    logger.info("Extracted %s: %d bytes, %s pages, %d chars", kind, len(data), doc.page_count, len(doc.text))
    return doc


# ---------------------------------------------------------------- PII stripping

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# Scheme/www URLs always; bare domains only when a path follows (github.com/user), so that
# technology names such as "ASP.NET" or "Node.js" are left alone.
_URL = re.compile(
    r"(?:https?://|www\.)[^\s<>\"')\]]+"
    r"|\b(?:[a-z0-9-]+\.)+(?:com|in|io|dev|me|app|net|org|co|tech|xyz|ai|design|site|page)/[^\s<>\"')\]]+"
)
_PHONE = re.compile(r"(?<![\w.])\+?\(?\d{1,5}\)?(?:[\s.-]?\(?\d{2,5}\)?){1,5}(?![\w.]*\d)")
_NOT_A_NAME = {"resume", "curriculum vitae", "cv", "profile", "contact", "summary", "bio data", "biodata"}


@dataclass
class StrippedText:
    text: str
    # placeholder -> original value; stays on the server, used by restore_pii
    mapping: dict[str, str] = field(default_factory=dict)


class _Replacer:
    def __init__(self) -> None:
        self.mapping: dict[str, str] = {}
        self._by_value: dict[tuple[str, str], str] = {}

    def placeholder(self, kind: str, value: str) -> str:
        key = (kind, value.lower())
        if key not in self._by_value:
            n = sum(1 for k, _ in self._by_value if k == kind) + 1
            self._by_value[key] = f"[{kind}_{n}]"
            self.mapping[self._by_value[key]] = value
        return self._by_value[key]


def _header_name(text: str) -> str | None:
    """The first non-empty line, if it looks like a person's name."""
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        words = line.split()
        looks_like_name = (
            2 <= len(words) <= 4
            and len(line) <= 40
            and all(re.fullmatch(r"[^\W\d_][^\W\d_.'-]*\.?", w) for w in words)
            and line.lower() not in _NOT_A_NAME
        )
        return line if looks_like_name else None
    return None


def strip_pii(
    text: str, known_names: tuple[str, ...] | list[str] = (), *, header_name: bool = True
) -> StrippedText:
    """Replace emails, phone numbers, URLs and the person's name with placeholders.

    `header_name` also treats a name-like first line as the person's name: right for a resume, wrong
    for a README or a project description, whose first line is a title.
    """
    rep = _Replacer()

    def sub_trimmed(kind: str):
        def _sub(match: re.Match) -> str:
            value = match.group(0)
            trailing = value[len(value.rstrip(".,;:")) :]
            value = value[: len(value) - len(trailing)] if trailing else value
            return rep.placeholder(kind, value) + trailing

        return _sub

    out = _EMAIL.sub(sub_trimmed("EMAIL"), text)
    out = _URL.sub(sub_trimmed("URL"), out)

    def _phone(match: re.Match) -> str:
        value = match.group(0)
        digits = sum(c.isdigit() for c in value)
        if not 10 <= digits <= 13:  # years, ranges, metrics and ids are not phone numbers
            return value
        return rep.placeholder("PHONE", value.strip())

    out = _PHONE.sub(_phone, out)

    names = [
        n.strip() for n in (*known_names, (_header_name(out) if header_name else "") or "") if n and n.strip()
    ]
    if names:
        rep.mapping["[NAME]"] = names[0]
        parts: set[str] = set()
        for name in names:
            parts.add(name)
            parts.update(w for w in re.split(r"[\s.]+", name) if len(w) >= 3)
        # longest first, so the full name goes before its parts
        for part in sorted(parts, key=len, reverse=True):
            out = re.sub(rf"(?<!\w){re.escape(part)}(?!\w)", "[NAME]", out, flags=re.IGNORECASE)
        out = re.sub(r"\[NAME\](?:[\s.]+\[NAME\])+", "[NAME]", out)
    return StrippedText(out, rep.mapping)


def restore_pii(value: Any, mapping: dict[str, str]) -> Any:
    """Put original values back into a string or any nested structure (done locally, after the LLM)."""
    if not mapping:
        return value
    if isinstance(value, str):
        for placeholder, original in mapping.items():
            value = value.replace(placeholder, original)
        return value
    if isinstance(value, BaseModel):
        return type(value).model_validate(restore_pii(value.model_dump(), mapping))
    if isinstance(value, list):
        return [restore_pii(v, mapping) for v in value]
    if isinstance(value, dict):
        return {k: restore_pii(v, mapping) for k, v in value.items()}
    return value
