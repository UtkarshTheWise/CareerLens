"""Regenerate the binary fixtures in this folder. All content is synthetic.

    uv run python tests/fixtures/make_fixtures.py

Writes resume.pdf (2 pages of text), resume.docx and no_text.pdf (a drawing, no text layer).
The PDF is assembled by hand so the test suite needs no PDF-writing dependency.
"""

from pathlib import Path

HERE = Path(__file__).resolve().parent

RESUME_PAGES = [
    [
        "Aarav Mehta",
        "aarav.mehta@example.com | +91 98765 43210 | github.com/careerlens-demo",
        "linkedin.com/in/aarav-mehta-demo | https://aarav-demo.example.dev",
        "",
        "EDUCATION",
        "B.Tech Computer Science and Engineering, Example Institute of Technology, 2023 - 2027",
        "CGPA 8.4",
        "",
        "SKILLS",
        "Python, FastAPI, SQL, Docker, pytest, React, Kubernetes, AWS",
        "",
        "EXPERIENCE",
        "Backend Intern, Northwind Labs, 2025-05 to 2025-07",
        "- Reduced API response time by 35% by adding Redis caching to two endpoints",
        "- Wrote 40 pytest tests for the booking service",
        "- Documented the deployment steps for new interns",
    ],
    [
        "PROJECTS",
        "campus-api",
        "REST API for campus event registration built with FastAPI and PostgreSQL.",
        "Handled 10,000 requests in a local load test. Packaged with Docker.",
        "https://github.com/careerlens-demo/campus-api",
        "",
        "weather-dashboard",
        "Web page that shows the weather for a searched city using React.",
        "https://github.com/careerlens-demo/weather-dashboard",
        "",
        "CERTIFICATIONS",
        "Intro to SQL, Kaggle",
    ],
]


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(pages: list[bytes]) -> bytes:
    """Assemble a minimal PDF from one content stream per page."""
    objects: list[bytes] = []
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, stream in enumerate(pages):
        content_ref = 5 + 2 * i
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_ref} 0 R >>"
            ).encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def text_stream(lines: list[str]) -> bytes:
    parts = ["BT", "/F1 11 Tf", "14 TL", "56 740 Td"]
    for line in lines:
        parts.append(f"({_escape(line)}) Tj T*")
    parts.append("ET")
    return "\n".join(parts).encode("cp1252")


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    for page in RESUME_PAGES:
        for line in page:
            if line:
                document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Languages"
    table.rows[0].cells[1].text = "English, Hindi"
    document.save(str(path))


def main() -> None:
    (HERE / "resume.pdf").write_bytes(build_pdf([text_stream(p) for p in RESUME_PAGES]))
    (HERE / "no_text.pdf").write_bytes(build_pdf([b"0.8 g\n72 500 300 200 re\nf"]))
    make_docx(HERE / "resume.docx")
    print("wrote resume.pdf, no_text.pdf, resume.docx")


if __name__ == "__main__":
    main()
